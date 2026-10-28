"""Builds an anonymised, synthetic resume PDF/DOCX for tests that must run anywhere (CI).

The person and every fact below are fictional. Layout mimics a typical Word-exported resume:
centred name, contact line with links, bold section headings, right-aligned dates/locations
and round bullets.
"""

from __future__ import annotations

import io

from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

NAME = "JORDAN AVERY LEE"
EMAIL = "jordan.lee@example.com"
PHONE = "+1 555 010 0199"
LINKEDIN = "https://www.linkedin.com/in/jordan-avery-lee-example"

SECTIONS = [
    ("SKILLS", [
        ("skill", "Languages: Python, SQL, R"),
        ("skill", "Machine Learning: PyTorch, Scikit-learn, Pandas, XGBoost"),
        ("skill", "Cloud & Tools: AWS (S3, Lambda), Docker, Git, Tableau"),
    ]),
    ("EXPERIENCE", [
        ("role", ("DATA ANALYST", "Rochester, NY", "Lakeview Credit Union", "Jun 2022 - Aug 2024")),
        ("bullet", "Built weekly loan performance dashboards in Tableau using SQL views, cutting manual reporting time by 6 hours per week."),
        ("bullet", "Developed an XGBoost churn model in Python with Scikit-learn, improving retention campaign targeting accuracy by 12%."),
        ("bullet", "Collaborated with the lending team to define data quality checks for 3 core tables."),
        ("role", ("DATA INTERN", "Rochester, NY", "Genesee Health Network", "May 2021 - Aug 2021")),
        ("bullet", "Cleaned and merged patient scheduling data with Pandas to support a no-show analysis."),
    ]),
    ("PROJECTS", [
        ("project", "STREAMLIT FORECASTING APP (PERSONAL)"),
        ("bullet", "Built a demand forecasting prototype using PyTorch and Docker, deployed to AWS Lambda for a class demo."),
    ]),
    ("EDUCATION", [
        ("edu", ("Rochester Institute of Example", "Rochester, NY")),
        ("edu2", ("Bachelor of Science in Statistics", "May 2022")),
    ]),
    ("CERTIFICATIONS", [
        ("cert", ("Tableau Desktop Specialist", "Mar 2023")),
    ]),
]


def build_synthetic_resume_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=LETTER)
    width, height = LETTER
    left, right = 40, width - 40
    y = height - 50
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(width / 2, y, NAME)
    y -= 18
    c.setFont("Helvetica", 10)
    contact = f"{EMAIL} | {PHONE} | LinkedIn"
    c.drawCentredString(width / 2, y, contact)
    tw = c.stringWidth(contact, "Helvetica", 10)
    lx = width / 2 + tw / 2 - c.stringWidth("LinkedIn", "Helvetica", 10)
    c.linkURL(LINKEDIN, (lx, y - 2, lx + c.stringWidth("LinkedIn", "Helvetica", 10), y + 10), relative=0)
    y -= 26
    for heading, items in SECTIONS:
        c.setFont("Helvetica-Bold", 11)
        c.drawString(left, y, heading)
        c.line(left, y - 3, right, y - 3)
        y -= 17
        for kind, payload in items:
            if kind == "skill":
                cat, rest = payload.split(":", 1)
                c.setFont("Helvetica-Bold", 10)
                c.drawString(left, y, cat + ":")
                c.setFont("Helvetica", 10)
                c.drawString(left + c.stringWidth(cat + ": ", "Helvetica-Bold", 10), y, rest.strip())
                y -= 13
            elif kind == "role":
                title, loc, org, dates = payload
                c.setFont("Helvetica-Bold", 10)
                c.drawString(left, y, title)
                c.drawRightString(right, y, loc)
                y -= 13
                c.setFont("Helvetica", 10)
                c.drawString(left, y, org)
                c.drawRightString(right, y, dates)
                y -= 14
            elif kind == "project":
                c.setFont("Helvetica-Bold", 10)
                c.drawString(left, y, payload)
                y -= 14
            elif kind == "bullet":
                c.setFont("Helvetica", 10)
                c.drawString(left + 6, y, "•")
                words, line = payload.split(), ""
                for w in words:
                    trial = (line + " " + w).strip()
                    if c.stringWidth(trial, "Helvetica", 10) > right - left - 20:
                        c.drawString(left + 16, y, line)
                        y -= 12
                        line = w
                    else:
                        line = trial
                c.drawString(left + 16, y, line)
                y -= 14
            elif kind == "edu":
                inst, loc = payload
                c.setFont("Helvetica-Bold", 10)
                c.drawString(left, y, inst)
                c.drawRightString(right, y, loc)
                y -= 13
            elif kind == "edu2":
                deg, date = payload
                c.setFont("Helvetica", 10)
                c.drawString(left, y, deg)
                c.drawRightString(right, y, date)
                y -= 15
            elif kind == "cert":
                name, date = payload
                c.setFont("Helvetica", 10)
                c.drawString(left + 6, y, "•")
                c.drawString(left + 16, y, name)
                c.drawRightString(right, y, date)
                y -= 14
        y -= 8
    c.showPage()
    c.save()
    return buf.getvalue()


def build_synthetic_resume_docx() -> bytes:
    import docx

    d = docx.Document()
    p = d.add_paragraph()
    r = p.add_run(NAME)
    r.bold = True
    d.add_paragraph(f"{EMAIL} | {PHONE} | {LINKEDIN}")
    for heading, items in SECTIONS:
        hp = d.add_paragraph()
        hr = hp.add_run(heading)
        hr.bold = True
        for kind, payload in items:
            if kind == "skill":
                d.add_paragraph(payload)
            elif kind == "role":
                title, loc, org, dates = payload
                rp = d.add_paragraph()
                rr = rp.add_run(f"{title}\t{loc}")
                rr.bold = True
                d.add_paragraph(f"{org}\t{dates}")
            elif kind == "project":
                pp = d.add_paragraph()
                pp.add_run(payload).bold = True
            elif kind == "bullet":
                d.add_paragraph(payload, style="List Bullet")
            elif kind == "edu":
                ep = d.add_paragraph()
                ep.add_run(f"{payload[0]}\t{payload[1]}").bold = True
            elif kind == "edu2":
                d.add_paragraph(f"{payload[0]}\t{payload[1]}")
            elif kind == "cert":
                d.add_paragraph(f"{payload[0]}\t{payload[1]}", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
