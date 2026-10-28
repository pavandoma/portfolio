"""Turn PDF / DOCX / text documents into visual lines with layout hints.

PDF extraction uses PyMuPDF. Spans are regrouped into visual lines by baseline, because
Word-exported PDFs often put right-aligned text (locations, dates) into separate blocks,
and bullet glyphs into their own lines. Long runs of spaces inside a span are treated as
column breaks (that is how right-aligned dates are often typeset).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from app.core.text import BULLET_CHARS, clean_text, collapse_ws

_GAP_SPLIT_RE = re.compile(r"\s{3,}|\t+")


@dataclass
class Segment:
    text: str
    x0: float
    x1: float
    bold: bool = False
    size: float = 0.0


@dataclass
class DocLine:
    segments: list[Segment]
    page: int
    y: float
    x0: float
    size: float
    bold: bool
    is_bullet: bool = False
    links: list[str] = field(default_factory=list)
    index: int = 0
    region: str = "body"  # body | header | footer

    @property
    def text(self) -> str:
        return collapse_ws(" ".join(s.text for s in self.segments))

    @property
    def left(self) -> str:
        return self.segments[0].text if self.segments else ""

    @property
    def right(self) -> str | None:
        return self.segments[-1].text if len(self.segments) > 1 else None

    @property
    def main(self) -> str:
        """Everything except a right-aligned trailing segment."""
        if len(self.segments) > 1:
            return collapse_ws(" ".join(s.text for s in self.segments[:-1]))
        return self.text


@dataclass
class ExtractedDocument:
    kind: str
    lines: list[DocLine]
    page_count: int = 1
    page_width: float = 612.0
    page_height: float = 792.0
    links: list[dict] = field(default_factory=list)
    fonts: list[str] = field(default_factory=list)
    image_count: int = 0
    header_footer_text: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(line.text for line in self.lines)


def _is_bullet_glyph(text: str) -> bool:
    t = text.strip()
    return len(t) == 1 and t in BULLET_CHARS


def extract_pdf(content: bytes) -> ExtractedDocument:
    import pymupdf

    with pymupdf.open(stream=content, filetype="pdf") as doc:
        all_lines: list[DocLine] = []
        links: list[dict] = []
        fonts: set[str] = set()
        images = 0
        page_w, page_h = 612.0, 792.0
        for pno, page in enumerate(doc):
            page_w, page_h = page.rect.width, page.rect.height
            for f in page.get_fonts():
                fonts.add(re.sub(r"^[A-Z]{6}\+", "", f[3]))
            images += len(page.get_images())
            for ln in page.get_links():
                if ln.get("uri"):
                    links.append({"uri": ln["uri"], "rect": tuple(ln["from"]), "page": pno})
            spans = []
            data = page.get_text("dict", flags=pymupdf.TEXT_PRESERVE_WHITESPACE)
            for block in data["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block["lines"]:
                    for s in line["spans"]:
                        if not s["text"].strip():
                            continue
                        x0, y0, x1, y1 = s["bbox"]
                        spans.append(
                            {
                                "text": clean_text(s["text"]),
                                "x0": x0, "x1": x1, "y0": y0, "y1": y1,
                                "base": s["origin"][1],
                                "size": s["size"],
                                "bold": bool(s["flags"] & 16) or "bold" in s["font"].lower(),
                            }
                        )
            spans.sort(key=lambda s: (round(s["base"], 0), s["x0"]))
            rows: list[list[dict]] = []
            for s in spans:
                # Bullet glyphs are often taller (different font), so compare vertical centres.
                centre = (s["y0"] + s["y1"]) / 2
                for row in rows:
                    rc = sum((r["y0"] + r["y1"]) / 2 for r in row) / len(row)
                    if abs(rc - centre) <= max(3.0, s["size"] * 0.35):
                        row.append(s)
                        break
                else:
                    rows.append([s])
            rows.sort(key=lambda row: min(r["y0"] for r in row))
            pending_bullet = False
            for row in rows:
                row.sort(key=lambda r: r["x0"])
                segments: list[Segment] = []
                for s in row:
                    parts = [p for p in _GAP_SPLIT_RE.split(s["text"]) if p.strip()]
                    if not parts:
                        continue
                    width = s["x1"] - s["x0"]
                    for i, part in enumerate(parts):
                        # Approximate x for split parts: the trailing part of a gap-split span is right-aligned.
                        px0 = s["x0"] if i == 0 else s["x1"] - width * len(part) / max(1, len(s["text"]))
                        px1 = s["x1"] if i == len(parts) - 1 else s["x0"] + width * 0.5
                        new = Segment(part.strip(), px0, px1, s["bold"], s["size"])
                        if segments and i == 0 and s["x0"] - segments[-1].x1 < 12:
                            prev = segments[-1]
                            joiner = "" if prev.text.endswith(("/", "-")) else " "
                            segments[-1] = Segment(
                                collapse_ws(prev.text + joiner + new.text), prev.x0, new.x1,
                                prev.bold and new.bold, max(prev.size, new.size),
                            )
                        else:
                            segments.append(new)
                if not segments:
                    continue
                is_bullet = False
                if _is_bullet_glyph(segments[0].text):
                    is_bullet = True
                    segments = segments[1:]
                elif segments[0].text[:1] in BULLET_CHARS and segments[0].text[1:2] == " ":
                    is_bullet = True
                    segments[0] = Segment(segments[0].text[2:].strip(), segments[0].x0, segments[0].x1,
                                          segments[0].bold, segments[0].size)
                if not segments:
                    pending_bullet = True
                    continue
                if pending_bullet:
                    is_bullet = True
                    pending_bullet = False
                y = min(r["y0"] for r in row)
                region = "body"
                if y < page_h * 0.03:
                    region = "header"
                elif y > page_h * 0.97:
                    region = "footer"
                text_segs = [s for s in segments if s.text]
                bold_chars = sum(len(s.text) for s in text_segs if s.bold)
                total_chars = sum(len(s.text) for s in text_segs) or 1
                dl = DocLine(
                    segments=text_segs,
                    page=pno,
                    y=y,
                    x0=text_segs[0].x0,
                    size=max(s.size for s in text_segs),
                    bold=bold_chars / total_chars > 0.6,
                    is_bullet=is_bullet,
                    region=region,
                )
                row_y0, row_y1 = min(r["y0"] for r in row), max(r["y1"] for r in row)
                for lk in links:
                    if lk["page"] == pno and lk["rect"][1] < row_y1 and lk["rect"][3] > row_y0:
                        dl.links.append(lk["uri"])
                all_lines.append(dl)
        for i, line in enumerate(all_lines):
            line.index = i
        header_footer = [line.text for line in all_lines if line.region != "body"]
        return ExtractedDocument(
            kind="pdf",
            lines=all_lines,
            page_count=doc.page_count,
            page_width=page_w,
            page_height=page_h,
            links=links,
            fonts=sorted(fonts),
            image_count=images,
            header_footer_text=header_footer,
        )


def extract_docx(content: bytes) -> ExtractedDocument:
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(io.BytesIO(content))
    lines: list[DocLine] = []
    links: list[dict] = []
    rels = document.part.rels
    for rel in rels.values():
        if "hyperlink" in rel.reltype and rel.is_external:
            links.append({"uri": rel.target_ref, "rect": None, "page": 0})

    def para_line(p) -> DocLine | None:
        raw = clean_text(p.text)
        if not raw.strip():
            return None
        style = (p.style.name or "").lower() if p.style is not None else ""
        numbered = p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
        is_bullet = "list" in style or "bullet" in style or numbered
        parts = [x for x in re.split(r"\t+|\s{3,}", raw) if x.strip()]
        if parts and _is_bullet_glyph(parts[0]):
            is_bullet = True
            parts = parts[1:]
        elif parts and parts[0][:1] in BULLET_CHARS and parts[0][1:2] == " ":
            is_bullet = True
            parts[0] = parts[0][2:]
        runs = [r for r in p.runs if r.text.strip()]
        bold_chars = sum(len(r.text) for r in runs if r.bold or (r.bold is None and p.style is not None and p.style.font.bold))
        total = sum(len(r.text) for r in runs) or 1
        sizes = [r.font.size.pt for r in runs if r.font.size is not None]
        size = max(sizes) if sizes else (p.style.font.size.pt if p.style is not None and p.style.font.size else 11.0)
        segs = [Segment(x.strip(), float(i * 300), float(i * 300 + 200), bold_chars / total > 0.6, size) for i, x in enumerate(parts)]
        if not segs:
            return None
        hyperlinks = []
        for h in p._p.iter(qn("w:hyperlink")):
            rid = h.get(qn("r:id"))
            if rid and rid in rels:
                hyperlinks.append(rels[rid].target_ref)
        return DocLine(segments=segs, page=0, y=0.0, x0=0.0, size=size, bold=bold_chars / total > 0.6,
                       is_bullet=is_bullet, links=hyperlinks)

    for p in document.paragraphs:
        line = para_line(p)
        if line:
            lines.append(line)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    line = para_line(p)
                    if line:
                        lines.append(line)
    header_footer: list[str] = []
    for section in document.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                if p.text.strip():
                    header_footer.append(clean_text(p.text).strip())
    for i, line in enumerate(lines):
        line.index = i
    fonts = sorted({r.font.name for p in document.paragraphs for r in p.runs if r.font.name})
    images = len(document.inline_shapes)
    warnings = []
    if document.tables:
        warnings.append("Document uses tables; some parsers read table cells out of order.")
    return ExtractedDocument(kind="docx", lines=lines, page_count=0, links=links, fonts=fonts,
                             image_count=images, header_footer_text=header_footer, warnings=warnings)


def extract_text(content: bytes) -> ExtractedDocument:
    try:
        raw = content.decode("utf-8")
    except UnicodeDecodeError:
        raw = content.decode("cp1252", errors="replace")
    lines: list[DocLine] = []
    for raw_line in clean_text(raw).splitlines():
        if not raw_line.strip():
            continue
        stripped = raw_line.strip()
        is_bullet = False
        if stripped[:1] in BULLET_CHARS and stripped[1:2] in (" ", "\t"):
            is_bullet, stripped = True, stripped[2:].strip()
        parts = [p for p in re.split(r"\t+|\s{3,}", stripped) if p.strip()]
        segs = [Segment(p.strip(), float(i * 300), float(i * 300 + 200)) for i, p in enumerate(parts)]
        upper = stripped.isupper() and len(stripped) < 60
        lines.append(DocLine(segments=segs, page=0, y=0.0, x0=0.0, size=11.0, bold=upper, is_bullet=is_bullet,
                             index=len(lines)))
    return ExtractedDocument(kind="txt", lines=lines)


def extract_document(kind: str, content: bytes) -> ExtractedDocument:
    if kind == "pdf":
        return extract_pdf(content)
    if kind == "docx":
        return extract_docx(content)
    return extract_text(content)
