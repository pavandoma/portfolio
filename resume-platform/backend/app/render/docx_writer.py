"""Editable DOCX rendering with python-docx.

Uses Word's own "List Bullet" style for bullets, right-aligned tab stops for dates, and
keeps everything in the document body (no text boxes, tables, headers or footers), which is
what most ATS parsers handle best.
"""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas.resume import TailoredResume

WORD_FONTS = {"sans": "Calibri", "serif": "Times New Roman"}


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def _hyperlink(paragraph, url: str, text: str, font: str, size: float) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    r_pr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def render_docx(resume: TailoredResume) -> bytes:
    d = docx.Document()
    font = WORD_FONTS.get(resume.layout.font_family, "Calibri")
    size = resume.layout.body_font_size
    section = d.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    m = Inches(resume.layout.margin_in)
    section.left_margin = section.right_margin = m
    section.top_margin = section.bottom_margin = Inches(resume.layout.margin_in * 0.85)
    text_width = section.page_width - section.left_margin - section.right_margin

    normal = d.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    bullet_style = d.styles["List Bullet"]
    bullet_style.font.name = font
    bullet_style.font.size = Pt(size)
    bullet_style.paragraph_format.space_after = Pt(0.6)

    def para(text: str = "", *, bold=False, italic=False, align=None, size_pt=None, space_before=0.0):
        p = d.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size_pt:
                r.font.size = Pt(size_pt)
        return p

    def two_col(left: str, right: str | None, *, bold_left=False):
        p = d.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(left)
        r.bold = bold_left
        if right:
            p.add_run("\t" + right)
        return p

    def heading(text: str):
        p = para(text.upper(), bold=True, size_pt=size + 0.5, space_before=size * 0.75)
        _bottom_border(p)
        p.paragraph_format.space_after = Pt(2)

    def bullet(text: str):
        p = d.add_paragraph(style="List Bullet")
        p.add_run(text)

    para(resume.header.name.upper(), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=size + 6.5)
    para(resume.header.target_title, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=size + 1.5)
    cp = d.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links = resume.header.links + [""] * len(resume.header.contact)
    for i, (item, link) in enumerate(zip(resume.header.contact, links, strict=False)):
        if i:
            sep = cp.add_run("  |  ")
            sep.font.size = Pt(size - 0.5)
        if link:
            _hyperlink(cp, link, item, font, size - 0.5)
        else:
            r = cp.add_run(item)
            r.font.size = Pt(size - 0.5)

    for sec in resume.section_order:
        if sec == "summary" and resume.summary:
            para(" ".join(s.text for s in resume.summary), space_before=size * 0.5)
        elif sec == "skills" and resume.skills:
            heading("Skills")
            for line in resume.skills:
                p = d.add_paragraph()
                r = p.add_run(f"{line.category}: ")
                r.bold = True
                p.add_run(", ".join(line.items))
        elif sec == "experience" and resume.experience:
            heading("Professional Experience")
            for i, e in enumerate(resume.experience):
                p = two_col(e.title, e.location, bold_left=True)
                if i:
                    p.paragraph_format.space_before = Pt(size * 0.45)
                two_col(e.organization, e.date_range)
                for b in e.bullets:
                    bullet(b.text)
        elif sec == "projects" and resume.projects:
            all_academic = all(p.descriptor == "Academic project" for p in resume.projects)
            heading("Academic Projects" if all_academic else "Projects")
            for i, pr in enumerate(resume.projects):
                right = pr.date_range if all_academic else " | ".join(x for x in (pr.descriptor, pr.date_range) if x)
                p = two_col(pr.name, right or None, bold_left=True)
                if i:
                    p.paragraph_format.space_before = Pt(size * 0.3)
                for b in pr.bullets:
                    bullet(b.text)
        elif sec == "education" and resume.education:
            heading("Education")
            for i, ed in enumerate(resume.education):
                p = two_col(ed.institution, ed.location, bold_left=True)
                if i:
                    p.paragraph_format.space_before = Pt(size * 0.3)
                two_col(ed.degree + (f" (GPA: {ed.gpa})" if ed.gpa else ""), ed.date)
        elif sec == "certifications" and resume.certifications:
            heading("Certifications")
            for c in resume.certifications:
                p = d.add_paragraph(style="List Bullet")
                p.paragraph_format.tab_stops.add_tab_stop(text_width - Inches(0.25), WD_TAB_ALIGNMENT.RIGHT)
                p.add_run(c.name + (f" ({c.detail})" if c.detail else ""))
                if c.date:
                    p.add_run("\t" + c.date)
    for p in d.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
    core = d.core_properties
    core.author = resume.header.name
    core.title = f"{resume.header.name} - Resume"
    core.subject = resume.header.target_title
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def docx_to_pdf(docx_bytes: bytes, timeout: int = 60) -> bytes | None:
    """Convert with LibreOffice when it is installed; returns None if unavailable."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "resume.docx"
        src.write_bytes(docx_bytes)
        try:
            subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile",
                            "--convert-to", "pdf", "--outdir", tmp, str(src)],
                           check=True, capture_output=True, timeout=timeout)
        except (subprocess.SubprocessError, OSError):
            return None
        out = Path(tmp) / "resume.pdf"
        return out.read_bytes() if out.exists() else None
